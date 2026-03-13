# -*- coding: utf-8 -*-
"""
Admin interface for client management with 1C import button.
"""

import os
from django.contrib import admin, messages
from django.utils.translation import gettext_lazy as _
from django.template.response import TemplateResponse
from django.urls import path, reverse
from django.http import HttpResponseRedirect
from django.db.models import Count
from openpyxl import load_workbook
from .models import Client, ClientGroup


@admin.register(Client)
class ClientAdmin(admin.ModelAdmin):
    """Admin interface for Client model with 1C import."""

    list_display = ('name', 'code_1c', 'employee', 'address', 'get_groups', 'created_at', 'updated_at')
    list_filter = ('client_groups', 'employee')
    search_fields = ('name', 'address', 'code_1c', 'trading_point_name')
    filter_horizontal = ('client_groups',)
    list_per_page = 20
    change_list_template = 'admin/clients/client_import_change_list.html'

    def get_urls(self):
        """Add custom import URLs."""
        urls = super().get_urls()
        my_urls = [
            path('import-1c/', self.admin_site.admin_view(self.import_1c_view), name='clients_import_1c'),
            path('import-1c/confirm/', self.admin_site.admin_view(self.import_1c_confirm), name='clients_import_1c_confirm'),
        ]
        return my_urls + urls

    def import_1c_view(self, request):
        """Custom view for 1C import with confirmation."""
        if request.method == 'POST':
            import_file = request.FILES.get('import_file')
            if not import_file:
                self.message_user(request, 'Ошибка: Файл не выбран', level=messages.ERROR)
                return HttpResponseRedirect(reverse('admin:clients_client_changelist'))

            # Process and store in session
            wb = load_workbook(import_file)
            ws = wb.active

            updated = []
            not_found = []

            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
                code_1c = str(row[1]) if row[1] and str(row[1]).strip() else None
                name = row[2] if len(row) > 2 and row[2] else None

                if not name:
                    continue

                instance = None
                update_name = False

                # Step 1: Try to find by code_1c
                if code_1c:
                    instance = Client.objects.filter(code_1c=code_1c).first()
                    # If found by code_1c but name differs - mark for name update
                    if instance and name:
                        # Strip whitespace for comparison
                        db_name = instance.name.strip() if instance.name else ''
                        file_name = name.strip() if name else ''
                        if db_name != file_name:
                            update_name = True
                            print(f"Name mismatch: DB='{db_name}' vs File='{file_name}'")

                # Step 2: If not found, try to find by name
                if not instance and name:
                    instance = Client.objects.filter(name=name).first()

                # Store result - ONLY if there are actual changes
                if instance:
                    # Check if there are any changes
                    has_changes = False
                    has_code_change = False
                    
                    # Check code_1c change (only if file has code and it differs from DB)
                    if code_1c:
                        db_code = instance.code_1c.strip() if instance.code_1c else ''
                        if db_code != code_1c:
                            has_code_change = True
                            has_changes = True
                    
                    # Check name change (strip whitespace for comparison)
                    if name:
                        db_name = instance.name.strip() if instance.name else ''
                        file_name = name.strip() if name else ''
                        if db_name != file_name:
                            update_name = True
                            has_changes = True
                            print(f"Name mismatch: DB='{db_name}' vs File='{file_name}'")
                    
                    # Only add to list if there are changes
                    if has_changes:
                        updated_item = {
                            'row': row_idx,
                            'id': instance.id,
                            'name': name[:60],
                            'old_code': instance.code_1c,
                            'new_code': code_1c if has_code_change else None,
                            'old_name': instance.name if update_name else None,
                            'new_name': name if update_name else None,
                            'update_name': update_name,
                        }
                        updated.append(updated_item)
                        print(f"Added to update: id={instance.id}, update_name={update_name}, new_name='{name[:30] if update_name else None}'")
                else:
                    not_found.append({
                        'row': row_idx,
                        'name': name[:60],
                        'code_1c': code_1c
                    })

            # Store in session for confirmation
            request.session['import_data'] = {
                'updated': updated,
                'not_found': not_found,
            }

            # Show preview
            return TemplateResponse(request, 'admin/clients/client_import_preview.html', {
                **self.admin_site.each_context(request),
                'title': 'Предварительный просмотр импорта',
                'updated': updated,
                'updated_count': len(updated),
                'not_found': not_found,
                'not_found_count': len(not_found),
            })

        return TemplateResponse(request, 'admin/clients/client_import.html', {
            **self.admin_site.each_context(request),
            'title': 'Импорт клиентов из 1С',
        })

    def import_1c_confirm(self, request):
        """Confirm and apply import."""
        if request.method != 'POST':
            return HttpResponseRedirect(reverse('admin:clients_client_changelist'))

        import_data = request.session.get('import_data')
        if not import_data:
            self.message_user(request, 'Нет данных для импорта', level=messages.ERROR)
            return HttpResponseRedirect(reverse('admin:clients_client_changelist'))

        # Apply updates
        updated_count = 0
        name_updated_count = 0
        code_updated_count = 0

        for item in import_data['updated']:
            try:
                client = Client.objects.get(pk=item['id'])
                changed_fields = []

                # Update code_1c if provided and different
                if item.get('new_code'):
                    old_code = client.code_1c
                    client.code_1c = item['new_code']
                    changed_fields.append(f'code_1c')
                    code_updated_count += 1
                    print(f"Client {client.id}: code_1c {old_code} → {item['new_code']}")

                # Update name if marked for update
                if item.get('update_name') and item.get('new_name'):
                    old_name = client.name
                    client.name = item['new_name']
                    changed_fields.append(f'name')
                    name_updated_count += 1
                    print(f"Client {client.id}: name UPDATED '{old_name[:50]}' → '{item['new_name'][:50]}'")
                elif item.get('update_name'):
                    print(f"Client {client.id}: name flag set but new_name is empty")

                # Save if there are changes
                if changed_fields:
                    client.save()
                    updated_count += 1
                    print(f"Client {client.id} SAVED: {', '.join(changed_fields)}")
                else:
                    print(f"Client {client.id}: NO CHANGES to save (update_name={item.get('update_name')}, new_code={item.get('new_code')})")
            except Client.DoesNotExist:
                print(f"Client ID {item['id']} not found!")
                pass
            except Exception as e:
                print(f"Client {client.id} ERROR: {e}")
                pass

        # Clear session
        request.session['import_data'] = None

        msg = f'Импорт завершён! Обновлено клиентов: {updated_count}'
        if name_updated_count > 0:
            msg += f' (обновлено имён: {name_updated_count})'
        if code_updated_count > 0:
            msg += f' (обновлено code_1c: {code_updated_count})'

        self.message_user(request, msg, level=messages.SUCCESS)
        return HttpResponseRedirect(reverse('admin:clients_client_changelist'))

    def get_groups(self, obj):
        """Return comma-separated list of client groups."""
        return ', '.join([group.name for group in obj.client_groups.all()])
    get_groups.short_description = _('Группы')


@admin.register(ClientGroup)
class ClientGroupAdmin(admin.ModelAdmin):
    """Admin interface for ClientGroup model."""
    list_display = ('name', 'get_client_count', 'created_at')
    search_fields = ('name',)
    list_per_page = 20

    def get_client_count(self, obj):
        """Return number of clients in group."""
        return obj.client_set.count()
    get_client_count.short_description = _('Количество клиентов')
